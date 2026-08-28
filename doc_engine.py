"""
Moteur de génération de documents UM-ACEP.
Remplit les templates .docx à partir de données utilisateur.
"""
import html
import re
from pathlib import Path
from io import BytesIO
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE_DIR = Path(__file__).parent
TEMPLATES_DIR = BASE_DIR / "templates_docx"

# Configuration de chaque type de document :
# - template : nom du fichier .docx source
# - fields : liste des champs du formulaire (clé, libellé, type html)
# - row_map : {index de ligne du tableau -> clé de champ} pour la colonne "valeur"
# - date_in_label_row : {index de ligne -> clé de champ} quand la date doit être
#   insérée directement dans le libellé (ex: "... à la date du .......")
DOC_TYPES = {
    "solde": {
        "label": "Attestation de solde",
        "description": "Solde du compte à une date donnée.",
        "icon": "💰",
        "template": "ATTESTATION_DE_SOLDE_UM-ACEP.docx",
        "fields": [
            ("numero", "Numéro de l'attestation (ex: 0123)", "text"),
            ("date_document", "Date de délivrance (ex: 14 août 2026)", "text"),
            ("nom_prenom", "Nom et prénom(s) du titulaire", "text"),
            ("piece_identite", "Nature et n° de la pièce d'identité", "text"),
            ("numero_compte", "N° de compte", "text"),
            ("date_ouverture", "Date d'ouverture du compte", "text"),
            ("date_solde", "Date d'arrêté du solde", "text"),
            ("solde_montant", "Solde du compte (montant)", "text"),
            ("solde_lettres", "Solde en toutes lettres", "text"),
        ],
        "row_map": {1: "nom_prenom", 2: "piece_identite", 3: "numero_compte",
                    4: "date_ouverture", 5: "solde_montant", 6: "solde_lettres"},
        "date_in_label_row": {5: "date_solde"},
    },
    "capacite": {
        "label": "Attestation de capacité financière",
        "description": "Capacité financière du titulaire à une date donnée.",
        "icon": "📈",
        "template": "ATTESTATION_DE_CAPACITE_FINANCIERE_UM-ACEP.docx",
        "fields": [
            ("numero", "Numéro de l'attestation (ex: 0123)", "text"),
            ("date_document", "Date de délivrance (ex: 14 août 2026)", "text"),
            ("nom_prenom", "Nom et prénom(s) du titulaire", "text"),
            ("piece_identite", "Nature et n° de la pièce d'identité", "text"),
            ("numero_compte", "N° de compte", "text"),
            ("date_ouverture", "Date d'ouverture du compte", "text"),
            ("date_solde", "Date d'arrêté du solde", "text"),
            ("solde_montant", "Solde du compte (montant)", "text"),
            ("solde_lettres", "Solde en toutes lettres", "text"),
        ],
        "row_map": {1: "nom_prenom", 2: "piece_identite", 3: "numero_compte",
                    4: "date_ouverture", 5: "solde_montant", 6: "solde_lettres"},
        "date_in_label_row": {5: "date_solde"},
    },
    "non_engagement": {
        "label": "Attestation de non engagement",
        "description": "Prêt intégralement remboursé, aucun engagement en cours.",
        "icon": "✅",
        "template": "ATTESTATION_DE_NON_ENGAGEMENT_PRET_UM-ACEP.docx",
        "fields": [
            ("numero", "Numéro de l'attestation (ex: 0123)", "text"),
            ("date_document", "Date de délivrance (ex: 14 août 2026)", "text"),
            ("nom_prenom", "Nom et prénom(s) du titulaire", "text"),
            ("piece_identite", "Nature et n° de la pièce d'identité", "text"),
            ("numero_pret", "N° de prêt / contrat de crédit", "text"),
            ("date_octroi", "Date d'octroi du prêt", "text"),
            ("montant_initial", "Montant initial accordé", "text"),
            ("objet_pret", "Objet du prêt", "text"),
            ("date_cloture", "Date de remboursement intégral (clôture)", "text"),
        ],
        "row_map": {1: "nom_prenom", 2: "piece_identite", 4: "numero_pret",
                    5: "date_octroi", 6: "montant_initial", 7: "objet_pret", 8: "date_cloture"},
        "date_in_label_row": {},
    },
    "engagement": {
        "label": "Attestation d'engagement",
        "description": "Prêt actif en cours de remboursement, situation détaillée.",
        "icon": "📄",
        "template": "ATTESTATION_D_ENGAGEMENT_PRET_UM-ACEP.docx",
        "fields": [
            ("numero", "Numéro de l'attestation (ex: 0123)", "text"),
            ("date_document", "Date de délivrance (ex: 14 août 2026)", "text"),
            ("nom_prenom", "Nom et prénom(s) du titulaire", "text"),
            ("piece_identite", "Nature et n° de la pièce d'identité", "text"),
            ("numero_pret", "N° de prêt / contrat de crédit", "text"),
            ("date_octroi", "Date d'octroi du prêt", "text"),
            ("montant_initial", "Montant initial accordé", "text"),
            ("objet_pret", "Objet du prêt", "text"),
            ("duree_totale", "Durée totale du prêt", "text"),
            ("date_encours", "Date d'arrêté de l'encours", "text"),
            ("encours_restant", "Encours restant dû", "text"),
            ("echeance_mensuelle", "Montant de l'échéance mensuelle", "text"),
            ("date_echeance_finale", "Date d'échéance finale", "text"),
        ],
        "row_map": {1: "nom_prenom", 2: "piece_identite", 5: "numero_pret",
                    6: "date_octroi", 7: "montant_initial", 8: "objet_pret", 9: "duree_totale",
                    12: "encours_restant", 13: "echeance_mensuelle", 14: "date_echeance_finale"},
        "date_in_label_row": {12: "date_encours"},
    },
}


def _set_paragraph_text(paragraph, new_text):
    if paragraph.runs:
        paragraph.runs[0].text = new_text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(new_text)


def _set_cell_text(cell, text):
    from docx.shared import Pt

    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        # aligne la mise en forme sur celle des libellés du tableau (Arial 10pt)
        run = p.add_run(text)
        run.font.name = "Arial"
        run.font.size = Pt(10)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def _build_document(doc_type_key, data):
    """Remplit le template (sans le sauvegarder) et retourne (doc, cfg)."""
    if doc_type_key not in DOC_TYPES:
        raise ValueError(f"Type de document inconnu: {doc_type_key}")
    cfg = DOC_TYPES[doc_type_key]
    doc = Document(TEMPLATES_DIR / cfg["template"])

    # En-tête : numéro d'attestation et date de délivrance
    for p in doc.paragraphs:
        text = "".join(r.text for r in p.runs)
        if "UM-ACEP / DIR" in text:
            new_text = re.sub(r"N°[^/]*/", f"N° {data.get('numero', '….')} /", text)
            _set_paragraph_text(p, new_text)
        elif text.strip().startswith("Dakar, le"):
            new_text = re.sub(r"(Dakar, le\s*)[.\s]*$", r"\g<1>" + data.get("date_document", ""), text)
            _set_paragraph_text(p, new_text)

    table = doc.tables[0]

    # Champs "valeur" dans la colonne de droite du tableau
    for row_idx, field_key in cfg["row_map"].items():
        cell = table.rows[row_idx].cells[1]
        _set_cell_text(cell, data.get(field_key, ""))

    # Dates insérées directement dans le libellé (ex: "... à la date du ...")
    for row_idx, field_key in cfg["date_in_label_row"].items():
        label_cell = table.rows[row_idx].cells[0]
        for p in label_cell.paragraphs:
            text = "".join(r.text for r in p.runs)
            if re.search(r"\.{3,}", text):
                new_text = re.sub(r"\.{3,}", data.get(field_key, ""), text)
                _set_paragraph_text(p, new_text)

    return doc, cfg


def generate_document(doc_type_key, data):
    """Remplit le template et retourne (BytesIO, nom_fichier_suggéré)."""
    doc, cfg = _build_document(doc_type_key, data)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    nom = data.get("nom_prenom", "document").strip().replace(" ", "_") or "document"
    filename = f"{cfg['template'].split('_UM-ACEP')[0]}_{nom}.docx"
    return buf, filename


# ---------------------------------------------------------------------------
# Rendu HTML fidèle au template (couleurs, gras, alignement, bordures, tableau)
# ---------------------------------------------------------------------------

_ALIGN_CSS = {
    WD_ALIGN_PARAGRAPH.LEFT: "left",
    WD_ALIGN_PARAGRAPH.CENTER: "center",
    WD_ALIGN_PARAGRAPH.RIGHT: "right",
    WD_ALIGN_PARAGRAPH.JUSTIFY: "justify",
}


def _iter_block_items(doc):
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, doc)
        elif child.tag == qn("w:tbl"):
            yield Table(child, doc)


def _run_css(run):
    parts = ["font-family:Arial, sans-serif"]
    if run.bold:
        parts.append("font-weight:bold")
    if run.italic:
        parts.append("font-style:italic")
    if run.underline:
        parts.append("text-decoration:underline")
    size = run.font.size
    if size:
        parts.append(f"font-size:{size.pt}pt")
    color = None
    try:
        if run.font.color is not None and run.font.color.rgb is not None:
            color = str(run.font.color.rgb)
    except Exception:
        color = None
    if color:
        parts.append(f"color:#{color}")
    return ";".join(parts)


def _is_title_paragraph(paragraph):
    """Détecte les titres d'attestation (ex: 'ATTESTATION DE SOLDE') : centrés,
    gras, grande taille de police."""
    if paragraph.alignment != WD_ALIGN_PARAGRAPH.CENTER:
        return False
    for run in paragraph.runs:
        if run.text.strip() and run.bold and run.font.size and run.font.size.pt >= 14:
            return True
    return False


# Deux niveaux de densité : "normal" (confortable) pour les documents courts,
# "compact" (légèrement resserré) pour les documents plus longs, afin que tout
# tienne sur une seule page sans être écrasé visuellement.
DENSITY_NORMAL = {
    "cell_pad": "6px 10px", "cell_pad_empty": "2px 10px",
    "table_gap": "22px", "para_margin": "4px 0", "title_top_margin": "34px",
    "title_bar_margin": "12px auto 44px", "line_height": "1.3",
}
DENSITY_COMPACT = {
    "cell_pad": "4px 10px", "cell_pad_empty": "1px 10px",
    "table_gap": "16px", "para_margin": "2px 0", "title_top_margin": "24px",
    "title_bar_margin": "9px auto 30px", "line_height": "1.15",
}


def _render_paragraph(paragraph, density, base_font_size="11pt", align_override=None):
    align = align_override or _ALIGN_CSS.get(paragraph.alignment, "left")
    spans = []
    for run in paragraph.runs:
        text = html.escape(run.text).replace("\t", "&emsp;")
        if text == "":
            continue
        spans.append(f'<span style="{_run_css(run)}">{text}</span>')
    is_empty = not spans
    if is_empty:
        # paragraphe vide (espacement uniquement dans le template) : hauteur réduite
        return '<p style="margin:0;line-height:0.5;font-size:6pt;">&nbsp;</p>'
    inner = "".join(spans)
    is_title = _is_title_paragraph(paragraph)
    margin = f'{density["title_top_margin"]} 0 4px' if is_title else density["para_margin"]
    para_html = (
        f'<p style="text-align:{align};margin:{margin};'
        f'font-size:{base_font_size};font-family:Arial, sans-serif;">{inner}</p>'
    )
    if is_title:
        para_html += (
            f'<div style="width:100%;height:3px;background:#02564A;'
            f'margin:{density["title_bar_margin"]};"></div>'
        )
    return para_html


def _cell_shading(tc):
    tcPr = tc.tcPr
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is not None:
        fill = shd.get(qn("w:fill"))
        if fill and fill.lower() != "auto":
            return fill
    return None


def _cell_gridspan(tc):
    tcPr = tc.tcPr
    if tcPr is None:
        return 1
    gs = tcPr.find(qn("w:gridSpan"))
    if gs is not None:
        val = gs.get(qn("w:val"))
        try:
            return int(val)
        except (TypeError, ValueError):
            return 1
    return 1


def _render_table(table, density):
    grid = table._tbl.find(qn("w:tblGrid"))
    widths = []
    if grid is not None:
        for col in grid.findall(qn("w:gridCol")):
            w = col.get(qn("w:w"))
            widths.append(int(w) if w else 0)
    total = sum(widths) or 1

    colgroup = ""
    if widths:
        cols = "".join(f'<col style="width:{(w / total) * 100:.2f}%">' for w in widths)
        colgroup = f"<colgroup>{cols}</colgroup>"

    # Bloc signature (ex: "Le Directeur Général") : pas de cadre, texte simple.
    is_signature = len(table.rows) == 1 and any(
        "directeur" in cell.text.lower() for cell in table.rows[0].cells
    )

    rows_html = []
    for row in table.rows:
        cells_html = []
        seen = set()
        for cell in row.cells:
            tc_id = id(cell._tc)
            if tc_id in seen:
                continue  # cellule déjà rendue (fusion horizontale)
            seen.add(tc_id)
            span = _cell_gridspan(cell._tc)
            fill = _cell_shading(cell._tc)
            is_empty = not any(p.text.strip() for p in cell.paragraphs)
            pad = density["cell_pad_empty"] if is_empty else density["cell_pad"]
            if is_signature:
                style = ["border:none", f"padding:{pad}", "vertical-align:top"]
            else:
                style = ["border:1px solid #D9D9D9", f"padding:{pad}", "vertical-align:top"]
            if fill:
                style.append(f"background:#{fill}")
            # les libellés de section (fond vert) sont centrés
            align_override = "center" if fill else None
            colspan_attr = f' colspan="{span}"' if span > 1 else ""
            inner = "".join(
                _render_paragraph(p, density, align_override=align_override)
                for p in cell.paragraphs
            )
            cells_html.append(f'<td style="{";".join(style)}"{colspan_attr}>{inner}</td>')
        rows_html.append(f"<tr>{''.join(cells_html)}</tr>")

    inner_table = (
        '<table style="border-collapse:collapse;width:100%;margin:0;">'
        + colgroup + "".join(rows_html) + "</table>"
    )

    # espace généreux avant/après chaque tableau
    table_margin = f'{density["table_gap"]} 0'

    if is_signature:
        return f'<div style="margin:{table_margin};">{inner_table}</div>'

    # Le cadre extérieur est porté par un <div> (et non par le <table> lui-même)
    # pour éviter un bug d'impression où la bordure droite d'un tableau à
    # border-collapse disparaît en bord de page.
    return (
        f'<div style="border:1.5px solid #1F3864;margin:{table_margin};'
        'overflow:hidden;">' + inner_table + "</div>"
    )


def render_docx_to_html(doc):
    """Reconstruit le document en HTML/CSS en respectant fidèlement la mise en
    forme du template : alignement, gras, couleurs, fond vert des en-têtes,
    bordures et largeurs de colonnes du tableau. Les documents comportant
    beaucoup de lignes utilisent un espacement légèrement plus resserré afin
    de toujours tenir sur une seule page à l'impression."""
    total_rows = sum(len(t.rows) for t in doc.tables)
    density = DENSITY_COMPACT if total_rows > 12 else DENSITY_NORMAL

    parts = []
    for block in _iter_block_items(doc):
        if isinstance(block, Paragraph):
            parts.append(_render_paragraph(block, density))
        elif isinstance(block, Table):
            parts.append(_render_table(block, density))
    content = "".join(parts)
    return f'<div style="line-height:{density["line_height"]}">{content}</div>'


def generate_preview_html(doc_type_key, data):
    """Génère le document et le convertit en HTML fidèle pour affichage à
    l'écran (aperçu non téléchargeable, impression uniquement)."""
    doc, cfg = _build_document(doc_type_key, data)
    html_content = render_docx_to_html(doc)
    nom = data.get("nom_prenom", "document").strip().replace(" ", "_") or "document"
    filename = f"{cfg['template'].split('_UM-ACEP')[0]}_{nom}.docx"
    return html_content, filename
